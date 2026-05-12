from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# Titre
title = doc.add_heading('Produits JetBrains Installes & Projets Associes', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Date
p = doc.add_paragraph(f'Genere le : {datetime.date.today().strftime("%d/%m/%Y")}')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# Donnees
data = [
    ('IntelliJ IDEA 2025.3', 'Java / Kotlin', [
        'fraudexai_complet',
        'PROGET PI (PFA 4eme annee)',
        'RI',
        'PROGET PI (Checkpoint Wael)',
        'DeepLearningP1',
    ]),
    ('PhpStorm 2025.3', 'PHP / Web', [
        'CyberHunter',
        'GestionDeProjets',
        'kitab',
        'RO',
        'presentation',
        'untitled',
        'Complexite',
        'FreaudDATA',
        'AppROcc',
        'GestionDesProjets',
        'AIDAA',
        'fraudexai_complet',
    ]),
    ('PhpStorm 2026.1', 'PHP / Web', [
        'CyberHunter',
        'GestionDeProjets',
        'kitab',
        'RO',
        'presentation',
        'untitled',
        'Complexite',
        'FreaudDATA',
        'AppROcc',
        'GestionDesProjets',
        'fraudexai_complet',
        'AIDAA',
        'latexFraud',
    ]),
    ('PyCharm 2026.1', 'Python', [
        'PyCharmMiscProject',
        'PythonProject',
    ]),
    ('Rider 2025.3', 'C# / .NET', [
        '1stProject',
        '1stproject.NET',
        'ConsoleApp11',
        'Test1',
        '1erProjet',
        'DEVOIRDS1',
    ]),
    ('Rider 2026.1', 'C# / .NET', [
        '1stProject',
        '1stproject.NET',
        'ConsoleApp11',
        'Test1',
        '1erProjet',
        'DEVOIRDS1',
    ]),
    ('DataSpell 2026.1', 'Data Science / Python', [
        'workspace (projet par defaut)',
    ]),
]

# Tableau
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# En-tete
hdr = table.rows[0].cells
hdr[0].text = 'Produit JetBrains'
hdr[1].text = 'Langage Principal'
hdr[2].text = 'Projets'

for cell in hdr:
    for para in cell.paragraphs:
        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run(para.text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), '1F497D')
    cell._tc.get_or_add_tcPr().append(shading)

# Lignes
colors = ['DEEAF1', 'FFFFFF']
for i, (product, lang, projects) in enumerate(data):
    row = table.add_row().cells
    row[0].text = product
    row[1].text = lang
    row[2].text = '\n'.join('- ' + p for p in projects)
    bg = colors[i % 2]
    for cell in row:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), bg)
        cell._tc.get_or_add_tcPr().append(shading)

# Largeurs colonnes
for row in table.rows:
    row.cells[0].width = Inches(2.2)
    row.cells[1].width = Inches(1.8)
    row.cells[2].width = Inches(3.5)

out = r'C:\Users\MohamedAbidi\IdeaProjects\DeepLearningP1\JetBrains_Produits_Projets.docx'
doc.save(out)
print('Fichier cree:', out)

